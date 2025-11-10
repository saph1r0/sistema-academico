"""
Comando de gestión para cargar estudiantes desde el archivo Excel global
"""
import os
from django.core.management.base import BaseCommand
from servicios.servicioEstudiantes import StudentLoader


class Command(BaseCommand):
    help = 'Carga estudiantes desde el archivo Excel global (bdtotall.xlsx)'
    
    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default='EXELS/bdtotall.xlsx',
            help='Ruta al archivo Excel de estudiantes (por defecto: EXELS/bdtotall.xlsx)'
        )
        
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Ejecutar sin hacer cambios en la base de datos'
        )
    
    def handle(self, *args, **options):
        file_path = options['file']
        dry_run = options['dry_run']
        
        # Verificar que el archivo existe
        if not os.path.exists(file_path):
            self.stdout.write(
                self.style.ERROR(f'❌ Archivo no encontrado: {file_path}')
            )
            return
        
        self.stdout.write(
            self.style.SUCCESS(f'🚀 Iniciando carga de estudiantes...')
        )
        
        if dry_run:
            self.stdout.write(
                self.style.WARNING('⚠️  MODO DRY-RUN: No se harán cambios en la base de datos')
            )
        
        try:
            # Crear instancia del servicio
            loader = StudentLoader()
            
            # Procesar archivo
            if not dry_run:
                results = loader.load_from_excel(file_path)
            else:
                # En modo dry-run, solo validar el archivo
                results = self._dry_run_validation(file_path)
            
            # Mostrar resultados
            self._display_results(results)
            
        except Exception as e:
            self.stdout.write(
                self.style.ERROR(f'❌ Error durante la carga: {str(e)}')
            )
            raise
    
    def _dry_run_validation(self, file_path):
        """Validación en modo dry-run sin hacer cambios"""
        import openpyxl
        
        try:
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active
            
            total_rows = sheet.max_row - 1  # Excluir encabezado
            
            self.stdout.write(f"📋 Archivo: {file_path}")
            self.stdout.write(f"📊 Hoja: {sheet.title}")
            self.stdout.write(f"📈 Filas de datos: {total_rows}")
            
            # Validar algunas filas de muestra
            sample_rows = min(5, total_rows)
            self.stdout.write(f"🔍 Validando {sample_rows} filas de muestra...")
            
            valid_count = 0
            for row_num in range(2, min(7, sheet.max_row + 1)):
                student_code = sheet.cell(row=row_num, column=1).value
                first_name = sheet.cell(row=row_num, column=2).value
                last_name = sheet.cell(row=row_num, column=3).value
                email = sheet.cell(row=row_num, column=4).value
                
                if all([student_code, first_name, last_name, email]):
                    valid_count += 1
                    self.stdout.write(f"   ✅ Fila {row_num}: {first_name} {last_name} ({email})")
                else:
                    self.stdout.write(f"   ❌ Fila {row_num}: Datos incompletos")
            
            return {
                'created': 0,
                'skipped': 0,
                'errors': 0,
                'total_rows': total_rows,
                'valid_sample': valid_count,
                'error_messages': []
            }
            
        except Exception as e:
            return {
                'created': 0,
                'skipped': 0,
                'errors': 1,
                'error_messages': [str(e)]
            }
    
    def _display_results(self, results):
        """Muestra los resultados del proceso"""
        self.stdout.write('\n' + '='*60)
        self.stdout.write(self.style.SUCCESS('📊 RESUMEN DE CARGA DE ESTUDIANTES'))
        self.stdout.write('='*60)
        
        if 'total_rows' in results:  # Modo dry-run
            self.stdout.write(f"📈 Total de filas en Excel: {results['total_rows']}")
            self.stdout.write(f"✅ Filas válidas en muestra: {results['valid_sample']}")
        else:  # Modo normal
            self.stdout.write(f"✅ Estudiantes creados: {results['created']}")
            self.stdout.write(f"⏭️  Estudiantes omitidos: {results['skipped']}")
            self.stdout.write(f"❌ Errores: {results['errors']}")
        
        if results['error_messages']:
            self.stdout.write('\n🚨 ERRORES ENCONTRADOS:')
            for error in results['error_messages'][:5]:
                self.stdout.write(f"   • {error}")
            
            if len(results['error_messages']) > 5:
                remaining = len(results['error_messages']) - 5
                self.stdout.write(f"   ... y {remaining} errores más")
        
        self.stdout.write('='*60)
        
        if results.get('created', 0) > 0:
            self.stdout.write(
                self.style.SUCCESS(f'🎉 ¡Proceso completado exitosamente!')
            )
        elif results.get('errors', 0) > 0:
            self.stdout.write(
                self.style.ERROR('⚠️  Proceso completado con errores')
            )
        else:
            self.stdout.write(
                self.style.WARNING('ℹ️  No se realizaron cambios')
            )