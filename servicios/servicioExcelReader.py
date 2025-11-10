"""
Servicio para lectura de archivos Excel del sistema de asignación de cursos
"""
import os
import logging
from typing import Dict, List, Optional, Tuple
from dataclasses import dataclass
import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


# Configurar logging
logger = logging.getLogger(__name__)


@dataclass
class StudentData:
    """Estructura de datos para estudiante"""
    email: str
    first_name: str
    last_name: str
    student_code: str
    full_name: str = ""
    
    def __post_init__(self):
        """Construir nombre completo después de inicialización"""
        self.full_name = f"{self.first_name} {self.last_name}".strip()


@dataclass
class ExcelReadResult:
    """Resultado de lectura de Excel"""
    success: bool
    data: Optional[Dict] = None
    error_message: str = ""
    warnings: List[str] = None
    
    def __post_init__(self):
        if self.warnings is None:
            self.warnings = []


class ExcelReader:
    """
    Servicio para leer archivos Excel de estudiantes y cursos
    
    Maneja la lectura de:
    - bdtotall.xlsx: archivo principal con todos los estudiantes
    - alumnos_*.xlsx: archivos de cursos específicos con códigos de estudiantes
    """
    
    def __init__(self):
        self.logger = logging.getLogger(self.__class__.__name__)
    
    def read_students_file(self, file_path: str) -> ExcelReadResult:
        """
        Lee el archivo bdtotall.xlsx y extrae datos de estudiantes
        
        Args:
            file_path: Ruta al archivo bdtotall.xlsx
            
        Returns:
            ExcelReadResult con diccionario {student_code: StudentData}
        """
        self.logger.info(f"Leyendo archivo de estudiantes: {file_path}")
        
        # Validar archivo
        validation_result = self._validate_file(file_path)
        if not validation_result.success:
            return validation_result
        
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            sheet = workbook.active
            
            students_data = {}
            warnings = []
            processed_count = 0
            
            # Leer datos desde la fila 3 (fila 2 son encabezados)
            for row_num in range(3, sheet.max_row + 1):
                try:
                    # Extraer datos de las columnas específicas
                    student_code = self._get_cell_value(sheet, row_num, 2)  # Columna B (CUI)
                    last_name_1 = self._get_cell_value(sheet, row_num, 4)   # Columna D (APELLIDO PATERNO)
                    last_name_2 = self._get_cell_value(sheet, row_num, 5)   # Columna E (APELLIDO MATERNO)
                    first_name = self._get_cell_value(sheet, row_num, 6)    # Columna F (NOMBRES)
                    email = self._get_cell_value(sheet, row_num, 7)         # Columna G (CORREO)
                    
                    # Validar datos requeridos
                    if not student_code or not email:
                        if student_code or email or first_name or last_name_1:
                            warnings.append(f"Fila {row_num}: Datos incompletos (código: {student_code}, email: {email})")
                        continue
                    
                    # Normalizar código de estudiante
                    student_code = str(student_code).strip()
                    if '.' in student_code:
                        student_code = student_code.split('.')[0]  # Remover decimales si existen
                    
                    # Normalizar email
                    email = str(email).strip().lower()
                    
                    # Construir nombres
                    first_name = str(first_name or "").strip()
                    last_name_1 = str(last_name_1 or "").strip()
                    last_name_2 = str(last_name_2 or "").strip()
                    
                    # Combinar apellidos
                    last_name = f"{last_name_1} {last_name_2}".strip()
                    
                    # Crear objeto StudentData
                    student_data = StudentData(
                        email=email,
                        first_name=first_name,
                        last_name=last_name,
                        student_code=student_code
                    )
                    
                    # Verificar duplicados
                    if student_code in students_data:
                        warnings.append(f"Código duplicado encontrado: {student_code} (fila {row_num})")
                        continue
                    
                    students_data[student_code] = student_data
                    processed_count += 1
                    
                except Exception as e:
                    warnings.append(f"Error procesando fila {row_num}: {str(e)}")
                    continue
            
            workbook.close()
            
            self.logger.info(f"Archivo leído exitosamente. Estudiantes procesados: {processed_count}")
            if warnings:
                self.logger.warning(f"Se encontraron {len(warnings)} advertencias durante la lectura")
            
            return ExcelReadResult(
                success=True,
                data=students_data,
                warnings=warnings
            )
            
        except Exception as e:
            error_msg = f"Error leyendo archivo {file_path}: {str(e)}"
            self.logger.error(error_msg)
            return ExcelReadResult(
                success=False,
                error_message=error_msg
            )
    
    def read_course_file(self, file_path: str) -> ExcelReadResult:
        """
        Lee un archivo alumnos_*.xlsx y extrae códigos de estudiantes
        
        Args:
            file_path: Ruta al archivo de curso
            
        Returns:
            ExcelReadResult con lista de códigos de estudiantes
        """
        self.logger.info(f"Leyendo archivo de curso: {file_path}")
        
        # Validar archivo
        validation_result = self._validate_file(file_path)
        if not validation_result.success:
            return validation_result
        
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            sheet = workbook.active
            
            student_codes = []
            warnings = []
            processed_count = 0
            
            # Buscar códigos de estudiante en todo el archivo
            for row_num in range(1, sheet.max_row + 1):
                for col_num in range(1, sheet.max_column + 1):
                    try:
                        cell_value = self._get_cell_value(sheet, row_num, col_num)
                        
                        if cell_value:
                            # Convertir a string y limpiar
                            str_value = str(cell_value).strip()
                            
                            # Verificar si es un código de estudiante (8 dígitos)
                            if str_value.replace('.', '').isdigit():
                                # Remover decimales si existen
                                clean_code = str_value.split('.')[0]
                                
                                # Verificar longitud típica de código de estudiante
                                if len(clean_code) == 8:
                                    if clean_code not in student_codes:
                                        student_codes.append(clean_code)
                                        processed_count += 1
                    
                    except Exception as e:
                        # Ignorar errores de celdas individuales
                        continue
            
            workbook.close()
            
            self.logger.info(f"Archivo de curso leído exitosamente. Códigos encontrados: {processed_count}")
            
            return ExcelReadResult(
                success=True,
                data=student_codes,
                warnings=warnings
            )
            
        except Exception as e:
            error_msg = f"Error leyendo archivo de curso {file_path}: {str(e)}"
            self.logger.error(error_msg)
            return ExcelReadResult(
                success=False,
                error_message=error_msg
            )
    
    def _validate_file(self, file_path: str) -> ExcelReadResult:
        """
        Valida que el archivo existe y es accesible
        
        Args:
            file_path: Ruta al archivo
            
        Returns:
            ExcelReadResult indicando si la validación fue exitosa
        """
        if not file_path:
            return ExcelReadResult(
                success=False,
                error_message="Ruta de archivo no proporcionada"
            )
        
        if not os.path.exists(file_path):
            return ExcelReadResult(
                success=False,
                error_message=f"Archivo no encontrado: {file_path}"
            )
        
        if not os.path.isfile(file_path):
            return ExcelReadResult(
                success=False,
                error_message=f"La ruta no es un archivo: {file_path}"
            )
        
        # Verificar extensión
        if not file_path.lower().endswith(('.xlsx', '.xls')):
            return ExcelReadResult(
                success=False,
                error_message=f"Formato de archivo no soportado: {file_path}"
            )
        
        # Verificar que el archivo no esté vacío
        try:
            file_size = os.path.getsize(file_path)
            if file_size == 0:
                return ExcelReadResult(
                    success=False,
                    error_message=f"El archivo está vacío: {file_path}"
                )
        except OSError as e:
            return ExcelReadResult(
                success=False,
                error_message=f"Error accediendo al archivo {file_path}: {str(e)}"
            )
        
        # Intentar abrir el archivo para verificar que no esté corrupto
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            if not workbook.worksheets:
                workbook.close()
                return ExcelReadResult(
                    success=False,
                    error_message=f"El archivo no contiene hojas de trabajo: {file_path}"
                )
            workbook.close()
        except Exception as e:
            return ExcelReadResult(
                success=False,
                error_message=f"Archivo corrupto o no válido {file_path}: {str(e)}"
            )
        
        return ExcelReadResult(success=True)
    
    def _get_cell_value(self, sheet: Worksheet, row: int, col: int) -> Optional[str]:
        """
        Obtiene el valor de una celda de forma segura
        
        Args:
            sheet: Hoja de trabajo
            row: Número de fila (1-indexed)
            col: Número de columna (1-indexed)
            
        Returns:
            Valor de la celda como string o None si está vacía
        """
        try:
            cell = sheet.cell(row=row, column=col)
            value = cell.value
            
            if value is None:
                return None
            
            # Convertir a string y limpiar espacios
            return str(value).strip()
            
        except Exception:
            return None
    
    def get_file_info(self, file_path: str) -> Dict:
        """
        Obtiene información básica de un archivo Excel
        
        Args:
            file_path: Ruta al archivo
            
        Returns:
            Diccionario con información del archivo
        """
        info = {
            'file_path': file_path,
            'exists': False,
            'readable': False,
            'sheets': [],
            'dimensions': None,
            'error': None
        }
        
        try:
            if os.path.exists(file_path):
                info['exists'] = True
                
                workbook = openpyxl.load_workbook(file_path, data_only=True)
                info['readable'] = True
                
                # Información de hojas
                info['sheets'] = [sheet.title for sheet in workbook.worksheets]
                
                # Dimensiones de la hoja activa
                if workbook.active:
                    sheet = workbook.active
                    info['dimensions'] = {
                        'rows': sheet.max_row,
                        'columns': sheet.max_column
                    }
                
                workbook.close()
                
        except Exception as e:
            info['error'] = str(e)
        
        return info