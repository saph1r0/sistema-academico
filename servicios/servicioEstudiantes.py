"""
Servicio para cargar estudiantes desde archivos Excel
"""
import openpyxl
import uuid
from django.db import connection, transaction
from django.contrib.auth.hashers import make_password
import re


class StudentLoader:
    """Servicio para cargar estudiantes desde el archivo Excel global"""
    
    def __init__(self):
        self.created_count = 0
        self.skipped_count = 0
        self.errors = []
    
    def load_from_excel(self, file_path):
        """
        Carga estudiantes desde el archivo Excel bdtotall.xlsx
        
        Args:
            file_path (str): Ruta al archivo Excel
            
        Returns:
            dict: Estadísticas del proceso de carga
        """
        print(f"📚 Iniciando carga de estudiantes desde: {file_path}")
        
        try:
            # Abrir el archivo Excel
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active
            
            print(f"📋 Procesando hoja: {sheet.title}")
            print(f"📊 Total de filas: {sheet.max_row}")
            
            # Procesar cada fila (empezar desde fila 3, ya que fila 2 son encabezados)
            with transaction.atomic():
                for row_num in range(3, sheet.max_row + 1):  # Empezar desde fila 3
                    try:
                        self._process_student_row(sheet, row_num)
                    except Exception as e:
                        error_msg = f"Error en fila {row_num}: {str(e)}"
                        self.errors.append(error_msg)
                        print(f"❌ {error_msg}")
                        continue
            
            # Mostrar estadísticas
            self._print_statistics()
            
            return {
                'created': self.created_count,
                'skipped': self.skipped_count,
                'errors': len(self.errors),
                'error_messages': self.errors
            }
            
        except Exception as e:
            error_msg = f"Error al procesar archivo Excel: {str(e)}"
            print(f"❌ {error_msg}")
            self.errors.append(error_msg)
            return {
                'created': 0,
                'skipped': 0,
                'errors': 1,
                'error_messages': [error_msg]
            }
    
    def _process_student_row(self, sheet, row_num):
        """Procesa una fila individual del Excel"""
        
        # Leer datos según la estructura real del Excel:
        # A=ID, B=CUI, C=INICIO, D=APELLIDO PATERNO, E=APELLIDO MATERNO, F=NOMBRES, G=CORREO
        student_id = self._get_cell_value(sheet, row_num, 1)        # Columna A (ID)
        cui = self._get_cell_value(sheet, row_num, 2)              # Columna B (CUI - código)
        apellido_paterno = self._get_cell_value(sheet, row_num, 4) # Columna D
        apellido_materno = self._get_cell_value(sheet, row_num, 5) # Columna E
        nombres = self._get_cell_value(sheet, row_num, 6)          # Columna F
        email = self._get_cell_value(sheet, row_num, 7)            # Columna G
        
        # Construir nombre completo y apellidos
        first_name = nombres if nombres else ""
        last_name = f"{apellido_paterno} {apellido_materno}".strip()
        student_code = cui  # Usar CUI como código de estudiante
        
        # Validar datos requeridos
        if not all([student_code, first_name, email]):
            missing_fields = []
            if not student_code: missing_fields.append("CUI")
            if not first_name: missing_fields.append("nombres")
            if not email: missing_fields.append("email")
            
            raise ValueError(f"Faltan campos requeridos: {', '.join(missing_fields)}")
        
        # Validar formato de email
        if not self._validate_email(email):
            raise ValueError(f"Email inválido: {email}")
        
        # Verificar si el usuario ya existe
        if self._user_exists(email):
            print(f"⏭️  Usuario ya existe: {email}")
            self.skipped_count += 1
            return
        
        # Crear usuario y estudiante
        self._create_user_and_student(
            email=email,
            first_name=first_name,
            last_name=last_name,
            student_code=student_code
        )
        
        print(f"✅ Creado: {first_name} {last_name} ({email})")
        self.created_count += 1
    
    def _get_cell_value(self, sheet, row, col):
        """Obtiene el valor de una celda y lo limpia"""
        cell = sheet.cell(row=row, column=col)
        value = cell.value
        
        if value is None:
            return ""
        
        # Convertir a string y limpiar espacios
        return str(value).strip()
    
    def _validate_email(self, email):
        """Valida el formato del email"""
        pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
        return re.match(pattern, email) is not None
    
    def _user_exists(self, email):
        """Verifica si un usuario ya existe por email"""
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT COUNT(*) FROM users WHERE institutional_email = %s",
                [email]
            )
            return cursor.fetchone()[0] > 0
    
    def _create_user_and_student(self, email, first_name, last_name, student_code):
        """Crea un usuario y su registro de estudiante correspondiente"""
        
        # Generar password temporal
        temp_password = self._generate_temp_password()
        hashed_password = make_password(temp_password)
        
        with connection.cursor() as cursor:
            # Crear usuario
            user_id = str(uuid.uuid4())
            cursor.execute("""
                INSERT INTO users (
                    id, institutional_email, password, first_name, last_name, 
                    role, is_active, date_joined, created_at, updated_at
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, NOW(), NOW(), NOW())
            """, [
                user_id, email, hashed_password, first_name, last_name,
                'student', True
            ])
            
            # Crear registro de estudiante
            student_id = str(uuid.uuid4())
            cursor.execute("""
                INSERT INTO students (
                    id, user_id, student_code, academic_status, created_at
                ) VALUES (%s, %s, %s, %s, NOW())
            """, [
                student_id, user_id, student_code, 'active'
            ])
            
            print(f"🔑 Password temporal para {email}: {temp_password}")
    
    def _generate_temp_password(self):
        """Genera una contraseña temporal"""
        import random
        import string
        
        # Generar password de 8 caracteres
        chars = string.ascii_letters + string.digits
        return ''.join(random.choice(chars) for _ in range(8))
    
    def _print_statistics(self):
        """Imprime estadísticas del proceso"""
        print("\n" + "="*50)
        print("📊 ESTADÍSTICAS DE CARGA DE ESTUDIANTES")
        print("="*50)
        print(f"✅ Estudiantes creados: {self.created_count}")
        print(f"⏭️  Estudiantes omitidos: {self.skipped_count}")
        print(f"❌ Errores encontrados: {len(self.errors)}")
        
        if self.errors:
            print("\n🚨 ERRORES DETALLADOS:")
            for error in self.errors[:10]:  # Mostrar solo los primeros 10
                print(f"   • {error}")
            
            if len(self.errors) > 10:
                print(f"   ... y {len(self.errors) - 10} errores más")
        
        print("="*50)