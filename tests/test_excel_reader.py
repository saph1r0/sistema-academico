"""
Tests para el servicio ExcelReader
"""
import unittest
import os
import tempfile
from unittest.mock import patch, MagicMock
import openpyxl

from servicios.servicioExcelReader import ExcelReader, StudentData, ExcelReadResult


class TestExcelReader(unittest.TestCase):
    """Tests para la clase ExcelReader"""
    
    def setUp(self):
        """Configuración inicial para cada test"""
        self.reader = ExcelReader()
    
    def test_student_data_creation(self):
        """Test creación de StudentData"""
        student = StudentData(
            email="test@unsa.edu.pe",
            first_name="Juan",
            last_name="Pérez García",
            student_code="20230001"
        )
        
        self.assertEqual(student.email, "test@unsa.edu.pe")
        self.assertEqual(student.first_name, "Juan")
        self.assertEqual(student.last_name, "Pérez García")
        self.assertEqual(student.student_code, "20230001")
        self.assertEqual(student.full_name, "Juan Pérez García")
    
    def test_excel_read_result_creation(self):
        """Test creación de ExcelReadResult"""
        # Resultado exitoso
        result = ExcelReadResult(success=True, data={"test": "data"})
        self.assertTrue(result.success)
        self.assertEqual(result.data, {"test": "data"})
        self.assertEqual(result.error_message, "")
        self.assertEqual(result.warnings, [])
        
        # Resultado con error
        result = ExcelReadResult(success=False, error_message="Test error")
        self.assertFalse(result.success)
        self.assertEqual(result.error_message, "Test error")
    
    def test_validate_file_nonexistent(self):
        """Test validación de archivo inexistente"""
        result = self.reader._validate_file("archivo_inexistente.xlsx")
        
        self.assertFalse(result.success)
        self.assertIn("no encontrado", result.error_message)
    
    def test_validate_file_empty_path(self):
        """Test validación con ruta vacía"""
        result = self.reader._validate_file("")
        
        self.assertFalse(result.success)
        self.assertIn("no proporcionada", result.error_message)
    
    def test_validate_file_wrong_extension(self):
        """Test validación con extensión incorrecta"""
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tmp:
            tmp.write(b"test content")
            tmp_path = tmp.name
        
        try:
            result = self.reader._validate_file(tmp_path)
            self.assertFalse(result.success)
            self.assertIn("no soportado", result.error_message)
        finally:
            os.unlink(tmp_path)
    
    def test_validate_file_empty_file(self):
        """Test validación de archivo vacío"""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name
        
        try:
            result = self.reader._validate_file(tmp_path)
            self.assertFalse(result.success)
            self.assertIn("vacío", result.error_message)
        finally:
            os.unlink(tmp_path)
    
    def test_get_cell_value_normal(self):
        """Test obtención de valor de celda normal"""
        # Crear un workbook temporal
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.cell(row=1, column=1, value="Test Value")
        
        value = self.reader._get_cell_value(sheet, 1, 1)
        self.assertEqual(value, "Test Value")
        
        wb.close()
    
    def test_get_cell_value_empty(self):
        """Test obtención de valor de celda vacía"""
        wb = openpyxl.Workbook()
        sheet = wb.active
        
        value = self.reader._get_cell_value(sheet, 1, 1)
        self.assertIsNone(value)
        
        wb.close()
    
    def test_get_cell_value_numeric(self):
        """Test obtención de valor numérico"""
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.cell(row=1, column=1, value=12345)
        
        value = self.reader._get_cell_value(sheet, 1, 1)
        self.assertEqual(value, "12345")
        
        wb.close()
    
    def test_get_file_info_nonexistent(self):
        """Test información de archivo inexistente"""
        info = self.reader.get_file_info("archivo_inexistente.xlsx")
        
        self.assertEqual(info['file_path'], "archivo_inexistente.xlsx")
        self.assertFalse(info['exists'])
        self.assertFalse(info['readable'])
        self.assertEqual(info['sheets'], [])
        self.assertIsNone(info['dimensions'])
    
    def create_test_students_excel(self):
        """Crea un archivo Excel de prueba para estudiantes"""
        wb = openpyxl.Workbook()
        sheet = wb.active
        
        # Encabezados en fila 2
        headers = ["ID", "CUI", "INICIO", "APELLIDO PATERNO", "APELLIDO MATERNO", 
                  "NOMBRES", "CORREO", "CURSOS", "SEMESTRE", "CREDITOS"]
        for col, header in enumerate(headers, 1):
            sheet.cell(row=2, column=col, value=header)
        
        # Datos de prueba desde fila 3
        test_data = [
            [1, 20230001, 2023, "PÉREZ", "GARCÍA", "JUAN CARLOS", "jperez@unsa.edu.pe", None, 1, 24],
            [2, 20230002, 2023, "LÓPEZ", "MARTÍN", "MARÍA ELENA", "mlopez@unsa.edu.pe", None, 1, 24],
            [3, 20230003, 2023, "GONZÁLEZ", "RUIZ", "PEDRO LUIS", "pgonzalez@unsa.edu.pe", None, 1, 24],
        ]
        
        for row_idx, row_data in enumerate(test_data, 3):
            for col_idx, value in enumerate(row_data, 1):
                sheet.cell(row=row_idx, column=col_idx, value=value)
        
        return wb
    
    def create_test_course_excel(self):
        """Crea un archivo Excel de prueba para curso"""
        wb = openpyxl.Workbook()
        sheet = wb.active
        
        # Simular estructura de archivo de curso con códigos dispersos
        sheet.cell(row=1, column=1, value="UNIVERSIDAD NACIONAL DE SAN AGUSTÍN")
        sheet.cell(row=5, column=2, value=20230001)
        sheet.cell(row=8, column=3, value=20230002)
        sheet.cell(row=12, column=1, value=20230003)
        sheet.cell(row=15, column=4, value="No es código")
        sheet.cell(row=18, column=2, value=12345)  # Código muy corto
        
        return wb
    
    def test_read_students_file_success(self):
        """Test lectura exitosa de archivo de estudiantes"""
        wb = self.create_test_students_excel()
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            wb.save(tmp.name)
            tmp_path = tmp.name
        
        wb.close()
        
        try:
            result = self.reader.read_students_file(tmp_path)
            
            self.assertTrue(result.success)
            self.assertIsInstance(result.data, dict)
            self.assertEqual(len(result.data), 3)
            
            # Verificar datos específicos
            student = result.data.get("20230001")
            self.assertIsNotNone(student)
            self.assertEqual(student.email, "jperez@unsa.edu.pe")
            self.assertEqual(student.first_name, "JUAN CARLOS")
            self.assertEqual(student.last_name, "PÉREZ GARCÍA")
            self.assertEqual(student.full_name, "JUAN CARLOS PÉREZ GARCÍA")
            
        finally:
            os.unlink(tmp_path)
    
    def test_read_course_file_success(self):
        """Test lectura exitosa de archivo de curso"""
        wb = self.create_test_course_excel()
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            wb.save(tmp.name)
            tmp_path = tmp.name
        
        wb.close()
        
        try:
            result = self.reader.read_course_file(tmp_path)
            
            self.assertTrue(result.success)
            self.assertIsInstance(result.data, list)
            self.assertEqual(len(result.data), 3)
            
            # Verificar códigos encontrados
            expected_codes = ["20230001", "20230002", "20230003"]
            for code in expected_codes:
                self.assertIn(code, result.data)
            
        finally:
            os.unlink(tmp_path)
    
    def test_read_students_file_with_missing_data(self):
        """Test lectura de archivo con datos faltantes"""
        wb = openpyxl.Workbook()
        sheet = wb.active
        
        # Encabezados
        headers = ["ID", "CUI", "INICIO", "APELLIDO PATERNO", "APELLIDO MATERNO", 
                  "NOMBRES", "CORREO", "CURSOS", "SEMESTRE", "CREDITOS"]
        for col, header in enumerate(headers, 1):
            sheet.cell(row=2, column=col, value=header)
        
        # Datos con información faltante
        sheet.cell(row=3, column=2, value=20230001)  # CUI sin email
        sheet.cell(row=4, column=7, value="test@unsa.edu.pe")  # Email sin CUI
        sheet.cell(row=5, column=2, value=20230002)  # CUI completo
        sheet.cell(row=5, column=7, value="complete@unsa.edu.pe")
        sheet.cell(row=5, column=4, value="APELLIDO")
        sheet.cell(row=5, column=6, value="NOMBRE")
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            wb.save(tmp.name)
            tmp_path = tmp.name
        
        wb.close()
        
        try:
            result = self.reader.read_students_file(tmp_path)
            
            self.assertTrue(result.success)
            self.assertEqual(len(result.data), 1)  # Solo el registro completo
            self.assertTrue(len(result.warnings) > 0)  # Debe haber advertencias
            
        finally:
            os.unlink(tmp_path)


if __name__ == '__main__':
    unittest.main()