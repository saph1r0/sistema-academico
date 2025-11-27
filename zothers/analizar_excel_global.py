#!/usr/bin/env python3
"""
Script para analizar la estructura del archivo Excel global
"""
import openpyxl

def analizar_excel():
    """Analiza la estructura del archivo bdtotall.xlsx"""
    
    try:
        wb = openpyxl.load_workbook('EXELS/bdtotall.xlsx')
        sheet = wb.active
        
        print(f"📋 Hoja: {sheet.title}")
        print(f"📊 Filas: {sheet.max_row}")
        print(f"📈 Columnas: {sheet.max_column}")
        
        # Mostrar encabezados (fila 2 parece ser la de encabezados)
        print("\n🏷️  ENCABEZADOS (Fila 2):")
        headers = []
        for col in range(1, sheet.max_column + 1):
            header = sheet.cell(row=2, column=col).value
            headers.append(str(header) if header else f"Col{col}")
            print(f"  Columna {col}: {header}")
        
        # Mostrar algunas filas de datos
        print(f"\n📄 PRIMERAS 5 FILAS DE DATOS:")
        for row in range(3, min(8, sheet.max_row + 1)):
            print(f"\nFila {row}:")
            for col in range(1, min(11, sheet.max_column + 1)):
                cell_value = sheet.cell(row=row, column=col).value
                header = headers[col-1] if col-1 < len(headers) else f"Col{col}"
                print(f"  {header}: {cell_value}")
        
        # Buscar columnas que podrían ser email
        print(f"\n🔍 BUSCANDO EMAILS EN LAS PRIMERAS 10 FILAS:")
        for col in range(1, sheet.max_column + 1):
            header = headers[col-1] if col-1 < len(headers) else f"Col{col}"
            
            # Revisar algunas celdas de esta columna
            sample_values = []
            for row in range(3, min(13, sheet.max_row + 1)):
                value = sheet.cell(row=row, column=col).value
                if value and '@' in str(value):
                    sample_values.append(str(value))
            
            if sample_values:
                print(f"  Columna {col} ({header}) contiene emails:")
                for email in sample_values[:3]:
                    print(f"    - {email}")
        
        print(f"\n✅ Análisis completado")
        
    except Exception as e:
        print(f"❌ Error: {e}")

if __name__ == "__main__":
    analizar_excel()